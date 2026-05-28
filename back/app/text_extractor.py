# app/text_extractor.py
from pathlib import Path
import re

import pdfplumber
from docx import Document as DocxDocument
from google.cloud import vision
from pptx import Presentation

PDF_PREVIEW_MAX_CHARS = 240


def get_pdf_page_count(file_path: str) -> int:
    """Retorna o total de paginas de um PDF."""
    with pdfplumber.open(file_path) as pdf:
        return len(pdf.pages)


def _normalize_pdf_text(text: str) -> str:
    normalized = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    normalized = re.sub(r"[ \t]+", " ", normalized)
    normalized = re.sub(r"(\w)-\n(\w)", r"\1\2", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def _extract_pdf_page_text(page: pdfplumber.page.Page) -> str:
    words = page.extract_words(
        use_text_flow=True,
        keep_blank_chars=False,
        extra_attrs=["top", "x0"],
    )

    if words:
        grouped_lines: list[tuple[float, list[tuple[float, str]]]] = []
        line_tolerance = 3.0

        for word in words:
            text = str(word.get("text", "")).strip()
            if not text:
                continue

            top = float(word.get("top", 0.0))
            x0 = float(word.get("x0", 0.0))

            if grouped_lines and abs(grouped_lines[-1][0] - top) <= line_tolerance:
                grouped_lines[-1][1].append((x0, text))
            else:
                grouped_lines.append((top, [(x0, text)]))

        rendered_lines: list[str] = []
        for _, line_words in grouped_lines:
            ordered_words = [text for _, text in sorted(line_words, key=lambda item: item[0])]
            rendered_line = " ".join(ordered_words).strip()
            if rendered_line:
                rendered_lines.append(rendered_line)

        extracted = "\n".join(rendered_lines).strip()
        if extracted:
            return _normalize_pdf_text(extracted)

    fallback_text = page.extract_text() or ""
    return _normalize_pdf_text(fallback_text)


def extract_text_from_pdf(file_path: str, pages: list[int] | None = None) -> str:
    """Extrai texto de um arquivo PDF."""
    extracted_pages: list[str] = []

    with pdfplumber.open(file_path) as pdf:
        target_pages = pdf.pages
        if pages is not None:
            target_pages = [pdf.pages[page_number - 1] for page_number in pages]

        for page in target_pages:
            page_text = _extract_pdf_page_text(page)
            if page_text.strip():
                extracted_pages.append(page_text)

    return "\n\n".join(extracted_pages)


def get_pdf_page_previews(
    file_path: str,
    max_chars: int = PDF_PREVIEW_MAX_CHARS,
) -> list[dict[str, str | int | bool]]:
    """Retorna uma amostra curta de texto por pagina para orientar a selecao."""
    previews: list[dict[str, str | int | bool]] = []

    with pdfplumber.open(file_path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            page_text = _extract_pdf_page_text(page)
            normalized_text = " ".join(page_text.split())
            has_text = bool(normalized_text)

            preview_text = normalized_text[:max_chars].rstrip()
            if len(normalized_text) > max_chars:
                preview_text = f"{preview_text}..."

            previews.append(
                {
                    "page_number": page_number,
                    "preview_text": preview_text,
                    "has_text": has_text,
                }
            )

    return previews


def extract_text_from_image(file_path: str) -> str:
    """Usa o Google Cloud Vision para extrair texto de uma imagem."""
    client = vision.ImageAnnotatorClient()

    with open(file_path, "rb") as image_file:
        content = image_file.read()

    image = vision.Image(content=content)

    response = client.text_detection(image=image)
    texts = response.text_annotations

    if response.error.message:
        raise Exception(
            f"{response.error.message}\nPara mais detalhes, veja https://cloud.google.com/apis/design/errors"
        )

    # O primeiro texto retornado e o texto completo detectado na imagem.
    return texts[0].description if texts else ""


def extract_text_from_docx(file_path: str) -> str:
    """Extrai texto de um arquivo DOCX preservando a ordem dos paragrafos."""
    document = DocxDocument(file_path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return "\n".join(paragraphs)


def extract_text_from_pptx(file_path: str) -> str:
    """Extrai texto de um arquivo PPTX agrupando o conteudo por slide."""
    presentation = Presentation(file_path)
    slides_content: list[str] = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        slide_parts: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", None)
            if text and text.strip():
                slide_parts.append(text.strip())

        if slide_parts:
            slides_content.append(f"Slide {slide_index}\n" + "\n".join(slide_parts))

    return "\n\n".join(slides_content)


def extract_text_from_file(file_path: str, pdf_pages: list[int] | None = None) -> str:
    """Despacha a extracao de texto com base na extensao do arquivo."""
    suffix = Path(file_path).suffix.lower()

    if suffix == ".pdf":
        return extract_text_from_pdf(file_path, pages=pdf_pages)
    if suffix in {".png", ".jpg", ".jpeg"}:
        return extract_text_from_image(file_path)
    if suffix == ".docx":
        return extract_text_from_docx(file_path)
    if suffix == ".pptx":
        return extract_text_from_pptx(file_path)

    raise ValueError(f"Tipo de ficheiro nao suportado: {suffix}")
